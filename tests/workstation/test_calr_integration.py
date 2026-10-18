from pathlib import Path
from mpn_tolkning.workstation.core.models import VariantRecord
from mpn_tolkning.workstation.services.browser_review import BrowserReviewService

def test_calr_non_calr_is_not_applicable_without_opening_browser(tmp_path, monkeypatch):
    service = BrowserReviewService(request_delay_ms=0)
    monkeypatch.setattr(service, '_browser_api', lambda: (_ for _ in ()).throw(AssertionError('network')))
    variant = VariantRecord(Path('synthetic.tsv'), 2, 'SYNTHETIC', 'JAK2', 'NM_004972.4:c.1849G>T')
    result = service.search_variants([variant], ['CALR-ETdb'], tmp_path)
    assert result[service.variant_key(variant)][0].status == 'not_applicable'

def test_calr_is_in_review_patient_and_capture_contract():
    from mpn_tolkning.workstation.services.evidence_audit import SCREENSHOT_REQUIRED_DATABASES
    from mpn_tolkning.workstation.reports.excel_report import DEFAULT_DATABASE_COLUMNS
    from mpn_tolkning.workstation.reports.patient_excel import REPORT_DATABASES
    from mpn_tolkning.workstation.gui.app import MainWindow
    from mpn_tolkning.workstation.services.settings import AppSettings
    assert 'CALR-ETdb' in DEFAULT_DATABASE_COLUMNS
    assert 'CALR-ETdb' in REPORT_DATABASES
    assert 'CALR-ETdb' in SCREENSHOT_REQUIRED_DATABASES
    assert 'CALR-ETdb' in MainWindow.databases
    assert 'CALR-ETdb' in AppSettings().enabled_databases


class Page:
    def __init__(self, html, tmp_path, *, mismatch=False, blank=False):
        self.html, self.tmp_path = html, tmp_path
        self.mismatch, self.blank = mismatch, blank
        self.submitted = []
        self.url = ""
    def goto(self, url, **kwargs): self.url = url
    def locator(self, selector):
        page = self
        class Locator:
            def fill(self, value): page.submitted.append((selector, value))
            def click(self): pass
            def wait_for(self, **kwargs): pass
            def inner_text(self): return "c.1154_1155insTTGTC" if page.mismatch else "c.1099_1150del"
            def screenshot(self, path): page.screenshot(path=path)
        return Locator()
    def evaluate(self, js): return self.html
    def screenshot(self, *, path, **kwargs):
        from PIL import Image, ImageDraw
        image = Image.new("RGB", (700, 200), "white")
        if not self.blank:
            ImageDraw.Draw(image).text((20, 20), "CALR c.1099_1150del type 1", fill="black")
        image.save(path)


def calr():
    return VariantRecord(Path('synthetic.tsv'), 2, 'NEVER-SUBMIT-THIS', 'CALR',
                         'NM_004343.4:c.1099_1150del', hgvsp='p.Leu367ThrfsTer46')


def test_browser_capture_uses_only_cdna_and_embeds_both_images(tmp_path):
    from mpn_tolkning.workstation.services.calr_et import _search_page
    from mpn_tolkning.workstation.core.models import ProcessingResult
    from mpn_tolkning.workstation.reports.patient_excel import PatientExcelReportWriter
    from openpyxl import load_workbook
    html = (Path(__file__).parents[1] / 'fixtures/calr_et_type1.html').read_text()
    page = Page(html, tmp_path)
    variant = calr()
    item = _search_page(BrowserReviewService(), page, variant, tmp_path)
    assert item.status == 'found'
    assert page.submitted == [('input[name="nucleotid"]', 'c.1099_1150del')]
    assert len(item.raw['screenshots']) == 2
    assert all(variant.sample not in r['path'] for r in item.raw['screenshots'])
    result = ProcessingResult(Path('synthetic.tsv'), None, '2026-10-18', [variant], [])
    output = tmp_path / 'patient.xlsx'
    PatientExcelReportWriter().write_patient(result, variant.patient_id, [variant], output,
        {BrowserReviewService.variant_key(variant): [item]})
    wb = load_workbook(output)
    assert len(wb['CALR']._images) == 2
    assert wb['Oversikt']['K10'].value == 'CALR-ETdb'
    assert wb['Oversikt']['K11'].value == 'Type 1-like; Class A'
    assert wb['Oversikt'].auto_filter.ref.endswith('K11')
    assert wb.sheetnames == ['Oversikt', 'Vedlegg', 'Data', 'CALR']
    wb.close()


def test_calr_capture_rejects_blank_and_wrong_detail(tmp_path):
    from mpn_tolkning.workstation.services.calr_et import _search_page
    html = (Path(__file__).parents[1] / 'fixtures/calr_et_type1.html').read_text()
    for options, status in [({'blank': True}, 'partial_capture'), ({'mismatch': True}, 'identity_mismatch')]:
        item = _search_page(BrowserReviewService(), Page(html, tmp_path, **options), calr(), tmp_path)
        assert item.status == status
        assert not item.clinical_significance


def test_calr_resume_retries_missing_capture_and_unverified_identity(tmp_path):
    from mpn_tolkning.workstation.core.models import DatabaseEvidence
    from mpn_tolkning.workstation.services.evidence_audit import migrate_loaded_evidence, is_completed_evidence
    item = DatabaseEvidence('CALR-ETdb', 'found', raw={'identity_verification': {'accepted': True}})
    migrate_loaded_evidence(item, tmp_path)
    assert item.status == 'partial_capture'
    assert not is_completed_evidence(item)
    legacy = DatabaseEvidence('CALR-ETdb', 'found')
    migrate_loaded_evidence(legacy, tmp_path)
    assert legacy.status == 'verification_required'
    negative = DatabaseEvidence('CALR-ETdb', 'not_found')
    migrate_loaded_evidence(negative, tmp_path)
    assert not is_completed_evidence(negative)


def test_original_mpn_preferences_migrate(tmp_path, monkeypatch):
    import json
    from mpn_tolkning.workstation.services.settings import AppSettings
    config = tmp_path / 'config.json'
    config.write_text(json.dumps({'output_directory': 'my-reports', 'selected_providers': ['CALR-ETdb'],
                                 'request_delay_min_ms': 2000, 'request_delay_max_ms': 4000}))
    monkeypatch.setattr(AppSettings, 'config_path', classmethod(lambda cls: config))
    settings = AppSettings.load()
    assert settings.default_output_dir == 'my-reports'
    assert settings.enabled_databases == ['CALR-ETdb']
    assert (settings.browser_delay_seconds, settings.browser_delay_max_seconds) == (2, 4)


def test_calr_workbook_roundtrip_preserves_evidence_and_recaptures_missing_image(tmp_path):
    from mpn_tolkning.workstation.core.models import ProcessingResult
    from mpn_tolkning.workstation.reports.excel_report import ExcelReportWriter
    from mpn_tolkning.workstation.services.calr_et import _search_page
    from mpn_tolkning.workstation.services.evidence_audit import write_evidence_audit
    from mpn_tolkning.workstation.services.processed_workbook import ProcessedWorkbookLoader
    variant = calr()
    variant.raw = {'Sample': variant.sample, 'Symbol': 'CALR', 'HGVSc': variant.hgvsc,
                   'HGVSp': variant.hgvsp, 'AF': 0.4, 'Depth': 1000, 'AO': 400,
                   'Type': 'deletion', 'Source': 'synthetic',
                   'Genomic Location': '', 'Ref/Alt Allele': ''}
    output = tmp_path / 'review.xlsx'
    directory = tmp_path / 'review_browser_evidence/calr-etdb'
    directory.mkdir(parents=True)
    html = (Path(__file__).parents[1] / 'fixtures/calr_et_type1.html').read_text()
    item = _search_page(BrowserReviewService(), Page(html, tmp_path), variant, directory)
    write_evidence_audit(directory, 'CALR-ETdb', variant, item)
    key = BrowserReviewService.variant_key(variant)
    result = ProcessingResult(Path('synthetic.tsv'), output, '2026-10-18', [variant], [])
    ExcelReportWriter().write(result, output, evidence={key: [item]})
    state = ProcessedWorkbookLoader().load(output)
    assert state.evidence[key][0].status == 'found'
    assert len(state.evidence[key][0].raw['screenshots']) == 2
    Path(item.raw['screenshots'][1]['path']).unlink()
    assert ProcessedWorkbookLoader().load(output).evidence[key][0].status == 'partial_capture'


def test_completed_calr_source_does_not_repeat_browser_search(tmp_path, monkeypatch):
    service = BrowserReviewService()
    monkeypatch.setattr(service, '_browser_api', lambda: (_ for _ in ()).throw(AssertionError('repeated browser')))
    variant = calr()
    key = service.variant_key(variant)
    assert service.search_variants([variant], ['CALR-ETdb'], tmp_path,
                                  completed_sources={(key, 'CALR-ETdb')}) == {key: []}
