from openpyxl import load_workbook
import pytest

from mpn_tolkning.core.models import EvidenceResult, EvidenceStatus
from mpn_tolkning.core.identity import stable_variant_key
from mpn_tolkning.reports.patient import write_patient_report


def test_calr_patient_report_uses_mpn_name_and_provider_order(tmp_path, calr_variant):
    evidence_set = {
        f"DIT-1001|{stable_variant_key(calr_variant)}|{provider}": EvidenceResult(provider, EvidenceStatus.FOUND, f"{provider} match")
        for provider in ("CALR-ETdb", "MTBP", "Franklin", "ClinVar", "OncoKB", "COSMIC")
    }
    path = write_patient_report(tmp_path, "DIT-1001", [calr_variant], evidence_set)
    assert path.name.startswith("DIT-1001-")
    assert path.name.endswith("_MPN_Tolkning.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames[:2] == ["Oversikt", "Vedlegg"]
    sheet = workbook[workbook.sheetnames[2]]
    captions = [cell.value for row in sheet.iter_rows() for cell in row if isinstance(cell.value, str)]
    indexes = [captions.index(name) for name in ("CALR-ETdb", "MTBP", "Franklin", "ClinVar", "OncoKB", "COSMIC")]
    assert indexes == sorted(indexes)


def test_report_contains_driver_context_and_review_disclaimer(tmp_path, jak2_variant):
    path = write_patient_report(tmp_path, "DIT-1002", [jak2_variant], [])
    workbook = load_workbook(path, data_only=True)
    overview = workbook["Oversikt"]
    text = " ".join(str(cell.value or "") for row in overview.iter_rows() for cell in row)
    assert "JAK2 V617F driver" in text
    assert "qualified" in text.lower()
    assert "research" in text.lower()


def test_unsafe_evidence_url_is_not_made_clickable(tmp_path, jak2_variant):
    prefix = f"DIT-1002|{stable_variant_key(jak2_variant)}|"
    evidence = {
        prefix + "Example": EvidenceResult("Example", EvidenceStatus.FOUND, "match", url="file:///secret"),
        prefix + "ClinVar": EvidenceResult("ClinVar", EvidenceStatus.FOUND, "match", url="https://evil.example/clinvar"),
    }
    path = write_patient_report(tmp_path, "DIT-1002", [jak2_variant], evidence)
    workbook = load_workbook(path)
    sheet = workbook[workbook.sheetnames[2]]
    hyperlinks = [cell.hyperlink for row in sheet.iter_rows() for cell in row if cell.hyperlink]
    assert hyperlinks == []
    values = {cell.value for row in sheet.iter_rows() for cell in row}
    assert "file:///secret" not in values
    assert "https://evil.example/clinvar" not in values


def test_allowlisted_provider_https_url_is_clickable(tmp_path, jak2_variant):
    evidence = {
        f"SYN|{stable_variant_key(jak2_variant)}|ClinVar": EvidenceResult(
            "ClinVar",
            EvidenceStatus.FOUND,
            "match",
            url="https://www.ncbi.nlm.nih.gov/clinvar/variation/123/",
        )
    }
    path = write_patient_report(tmp_path, "SYN", [jak2_variant], evidence)
    workbook = load_workbook(path)
    hyperlinks = [cell.hyperlink.target for row in workbook[workbook.sheetnames[2]].iter_rows() for cell in row if cell.hyperlink]
    assert hyperlinks == ["https://www.ncbi.nlm.nih.gov/clinvar/variation/123/"]


def test_keyed_evidence_is_scoped_to_its_variant_sheet(tmp_path, calr_variant, jak2_variant):
    evidence = {
        f"SYN|{stable_variant_key(calr_variant)}|CALR-ETdb": EvidenceResult(
            "CALR-ETdb", EvidenceStatus.FOUND, "CALR-only evidence"
        ),
        f"SYN|{stable_variant_key(jak2_variant)}|ClinVar": EvidenceResult(
            "ClinVar", EvidenceStatus.FOUND, "JAK2-only evidence"
        ),
    }
    path = write_patient_report(tmp_path, "SYN", [calr_variant, jak2_variant], evidence)
    workbook = load_workbook(path, data_only=True)
    calr_text = " ".join(str(cell.value or "") for row in workbook[workbook.sheetnames[2]].iter_rows() for cell in row)
    jak2_text = " ".join(str(cell.value or "") for row in workbook[workbook.sheetnames[3]].iter_rows() for cell in row)
    assert "CALR-only evidence" in calr_text and "JAK2-only evidence" not in calr_text
    assert "JAK2-only evidence" in jak2_text and "CALR-only evidence" not in jak2_text


def test_unkeyed_evidence_is_rejected_at_report_boundary(tmp_path, calr_variant, jak2_variant):
    evidence = [
        EvidenceResult(
            "ClinVar",
            EvidenceStatus.FOUND,
            "JAK2-only evidence",
            expected_identity=jak2_variant.hgvsc,
            observed_identity=jak2_variant.hgvsc,
        )
    ]
    with pytest.raises(ValueError, match="keyed by patient, variant, and provider"):
        write_patient_report(tmp_path, "SYN", [calr_variant, jak2_variant], evidence)
