import pytest
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QApplication

from mpn_tolkning.gui.app import MainWindow, Palette, trusted_calr_detail_url
from mpn_tolkning.core.models import EvidenceResult, EvidenceStatus


@pytest.fixture(scope="module")
def qt_app():
    return QApplication.instance() or QApplication([])


def test_workbench_has_approved_navigation(qt_app):
    window = MainWindow()
    assert [button.text() for button in window.nav_buttons] == [
        "Import", "Driver Review", "Evidence", "Reports", "Settings"
    ]
    assert window.page_stack.count() == 5


def test_driver_table_and_filter_have_accessible_names(qt_app):
    window = MainWindow()
    assert window.variant_table.accessibleName() == "MPN variant review table"
    assert window.review_filter.accessibleName() == "Filter MPN variants"


def test_window_respects_approved_minimum_size(qt_app):
    window = MainWindow()
    assert window.minimumWidth() == 1120
    assert window.minimumHeight() == 720


def test_sidebar_and_primary_actions_meet_target_sizes(qt_app):
    window = MainWindow()
    assert 190 <= window.sidebar.minimumWidth() <= 226
    assert window.import_button.minimumHeight() >= 44
    assert window.start_search_button.minimumHeight() >= 44


def test_imported_variants_populate_metrics_table_and_filter(qt_app, sample_variants):
    window = MainWindow()
    window.set_variants(sample_variants)
    assert window.variant_table.rowCount() == 3
    assert window.metric_values["variants"].text() == "3"
    assert window.metric_values["samples"].text() == "3"

    window.review_filter.setText("CALR")
    QApplication.processEvents()
    visible = sum(not window.variant_table.isRowHidden(row) for row in range(3))
    assert visible == 1


def test_table_selection_updates_variant_lens(qt_app, sample_variants):
    window = MainWindow()
    window.set_variants(sample_variants)
    window.variant_table.selectRow(1)
    QApplication.processEvents()
    assert "JAK2" in window.variant_title.text()
    assert "V617F" in window.driver_context_label.text()


def test_theme_contains_approved_semantic_palette(qt_app):
    window = MainWindow()
    style = window.styleSheet().upper()
    for color in (Palette.DEEP_TEAL, Palette.TEAL, Palette.AMBER, Palette.BACKGROUND):
        assert color.upper() in style


def test_navigation_changes_page_and_marks_current_button(qt_app):
    window = MainWindow()
    window.nav_buttons[2].click()
    assert window.page_stack.currentIndex() == 2
    assert window.nav_buttons[2].property("current") is True
    assert window.nav_buttons[2].accessibleName() == "Evidence page, current"


def test_calr_capture_accepts_only_the_expected_https_detail_endpoint():
    assert trusted_calr_detail_url(
        "https://www.dsimb.inserm.fr/CALR-ET/php/variant_page.php?num_var=var20"
    )
    assert not trusted_calr_detail_url("http://www.dsimb.inserm.fr/CALR-ET/php/variant_page.php?num_var=var20")
    assert not trusted_calr_detail_url("https://evil.example/CALR-ET/php/variant_page.php?num_var=var20")
    assert not trusted_calr_detail_url("https://www.dsimb.inserm.fr/CALR-ET/php/search_page.php")


def test_search_worker_resumes_existing_audit_results(qt_app, sample_variants):
    from mpn_tolkning.gui.workers import SearchWorker
    from mpn_tolkning.services.queue import EvidenceQueue

    class Provider:
        name = "ClinVar"
        calls = 0

        def search(self, variant, context):
            self.calls += 1
            return EvidenceResult(self.name, EvidenceStatus.FOUND, "fresh")

    provider = Provider()
    queue = EvidenceQueue([provider])
    key = queue.result_key(sample_variants[0], "ClinVar")
    existing = {key: EvidenceResult("ClinVar", EvidenceStatus.FOUND, "cached")}
    received = []
    worker = SearchWorker(queue, sample_variants[:1], existing_results=existing)
    worker.completed.connect(received.append)
    worker.run()
    assert received[0][key].summary == "cached"
    assert provider.calls == 0


def test_optional_identity_sources_are_available_but_disabled_by_default(qt_app, monkeypatch, tmp_path):
    monkeypatch.setenv("USERPROFILE", str(tmp_path))
    window = MainWindow()

    assert "ClinGen Allele Registry" in window.provider_checks
    assert "CIViC" in window.provider_checks
    assert "MetaKB" in window.provider_checks
    assert not window.provider_checks["CIViC"].isChecked()
    assert set(window.sign_in_buttons) == {"MTBP", "Franklin", "OncoKB", "COSMIC"}


def test_close_waits_for_active_workers(qt_app):
    class Event:
        ignored = False
        accepted = False

        def ignore(self):
            self.ignored = True

        def accept(self):
            self.accepted = True

    class Worker:
        stopped = False

        def stop(self):
            self.stopped = True

    window = MainWindow()
    worker = Worker()
    window._search_worker = worker
    window._threads.append(object())
    event = Event()

    window.closeEvent(event)

    assert event.ignored
    assert not event.accepted
    assert worker.stopped


def test_import_worker_builds_review_workbook_off_the_gui_thread(qt_app, tmp_path):
    from mpn_tolkning.gui.workers import ImportWorker

    source = tmp_path / "archer.tsv"
    source.write_text(
        "Sample\tGene\tHGVSc\tAF\nSYN\tJAK2\tNM_004972.4:c.1849G>T\t12.5\n",
        encoding="utf-8",
    )
    received = []
    worker = ImportWorker(source, tmp_path / "out", [])
    worker.completed.connect(received.append)

    worker.run()

    variants, workbook = received[0]
    assert variants[0].gene == "JAK2"
    assert workbook == tmp_path / "out" / "MPN_Variant_Review.xlsx"
    assert workbook.exists()


def test_reimport_preserves_human_review_selections(qt_app, tmp_path):
    from openpyxl import load_workbook
    from mpn_tolkning.gui.workers import ImportWorker

    source = tmp_path / "archer.tsv"
    source.write_text("Sample\tGene\tHGVSc\tAF\nSYN\tJAK2\tNM_004972.4:c.1849G>T\t12.5\n", encoding="utf-8")
    output = tmp_path / "out"
    first = []
    worker = ImportWorker(source, output, [])
    worker.completed.connect(first.append)
    worker.run()
    workbook_path = first[0][1]
    workbook = load_workbook(workbook_path)
    sheet = workbook["All Variants"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("Skip Database Search (X)") + 1, "X")
    workbook.save(workbook_path)

    second = []
    worker = ImportWorker(source, output, [])
    worker.completed.connect(second.append)
    worker.run()

    assert second[0][1] == workbook_path
    preserved = load_workbook(workbook_path, data_only=True)["All Variants"]
    assert preserved.cell(2, headers.index("Skip Database Search (X)") + 1).value == "X"
