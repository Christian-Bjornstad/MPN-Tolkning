from openpyxl import load_workbook
import pytest

from mpn_tolkning.core.context import classify_driver_context
from mpn_tolkning.core.identity import stable_variant_key
from mpn_tolkning.reports.review import (
    ReviewWorkbookError,
    load_skip_keys,
    write_review_workbook,
)
from mpn_tolkning.core.models import VariantRecord


def _contexts(variants):
    return {
        stable_variant_key(variant): classify_driver_context(variant)
        for variant in variants
    }


def test_review_workbook_preserves_rows_and_adds_skip_control(tmp_path, sample_variants):
    path = write_review_workbook(
        tmp_path / "review.xlsx",
        sample_variants,
        _contexts(sample_variants),
        {},
    )
    workbook = load_workbook(path)

    assert workbook.sheetnames == ["All Variants", "Artifacts Removed"]
    headers = [cell.value for cell in workbook["All Variants"][1]]
    assert headers[:4] == ["Sample", "Gene", "HGVSc", "HGVSp"]
    assert "Skip Database Search (X)" in headers
    assert "Stable Variant Key" in headers
    assert "Application Sample" in headers
    stable_column = headers.index("Stable Variant Key") + 1
    assert workbook["All Variants"].column_dimensions[
        workbook["All Variants"].cell(1, stable_column).column_letter
    ].hidden
    application_sample_column = headers.index("Application Sample") + 1
    assert workbook["All Variants"].column_dimensions[
        workbook["All Variants"].cell(1, application_sample_column).column_letter
    ].hidden


def test_artifact_removed_sheet_excludes_configured_artifact(tmp_path, sample_variants):
    artifact_key = stable_variant_key(sample_variants[1])
    path = write_review_workbook(
        tmp_path / "review.xlsx",
        sample_variants,
        _contexts(sample_variants),
        {artifact_key: "Configured test artifact"},
    )
    workbook = load_workbook(path, read_only=True)

    assert workbook["All Variants"].max_row == 4
    assert workbook["Artifacts Removed"].max_row == 3


def test_skip_selection_round_trips_by_sample_and_stable_key(tmp_path, sample_variants):
    path = write_review_workbook(
        tmp_path / "review.xlsx",
        sample_variants,
        _contexts(sample_variants),
        {},
    )
    workbook = load_workbook(path)
    sheet = workbook["All Variants"]
    headers = [cell.value for cell in sheet[1]]
    skip_col = headers.index("Skip Database Search (X)") + 1
    sheet.cell(2, skip_col, "x")
    workbook.save(path)

    assert load_skip_keys(path) == {
        f"{sample_variants[0].sample}|{stable_variant_key(sample_variants[0])}"
    }


def test_duplicate_review_identity_fails_closed(tmp_path, sample_variants):
    path = write_review_workbook(
        tmp_path / "review.xlsx",
        [sample_variants[0], sample_variants[0]],
        _contexts(sample_variants),
        {},
    )

    with pytest.raises(ReviewWorkbookError, match="Duplicate review identity"):
        load_skip_keys(path)


def test_untrusted_tsv_values_are_not_written_as_excel_formulas(tmp_path):
    variant = VariantRecord(
        sample="=HYPERLINK(\"https://evil.example\")",
        gene="JAK2",
        hgvsc="NM_004972.4:c.1849G>T",
        vaf=10.0,
        original={"Sample": "=HYPERLINK(\"https://evil.example\")", "Gene": "JAK2", "HGVSc": "NM_004972.4:c.1849G>T", "AF": 10.0},
    )
    path = write_review_workbook(tmp_path / "review.xlsx", [variant], {}, {})
    workbook = load_workbook(path, data_only=False)
    cell = workbook["All Variants"]["A2"]
    assert cell.data_type == "s"
    assert cell.value.startswith("'")


def test_skip_reload_accepts_supported_sample_id_alias(tmp_path, sample_variants):
    from dataclasses import replace

    variant = replace(
        sample_variants[0],
        original={"Sample ID": "SYN-ALIAS", "Gene": "CALR", "HGVSc": sample_variants[0].hgvsc, "AF": 20.0},
        sample="SYN-ALIAS",
    )
    path = write_review_workbook(tmp_path / "review.xlsx", [variant], _contexts([variant]), {})
    workbook = load_workbook(path)
    sheet = workbook["All Variants"]
    headers = [cell.value for cell in sheet[1]]
    sheet.cell(2, headers.index("Skip Database Search (X)") + 1, "X")
    workbook.save(path)
    assert load_skip_keys(path) == {f"SYN-ALIAS|{stable_variant_key(variant)}"}


def test_review_reload_rejects_stale_or_missing_rows(tmp_path, sample_variants):
    path = write_review_workbook(tmp_path / "review.xlsx", sample_variants[:1], _contexts(sample_variants[:1]), {})

    with pytest.raises(ReviewWorkbookError, match="do not match the current import"):
        load_skip_keys(path, sample_variants)


def test_review_reload_releases_file_handle_for_immediate_rewrite(tmp_path, sample_variants):
    path = write_review_workbook(tmp_path / "review.xlsx", sample_variants, _contexts(sample_variants), {})
    load_skip_keys(path, sample_variants)

    rewritten = write_review_workbook(path, sample_variants, _contexts(sample_variants), {})
    assert rewritten == path
