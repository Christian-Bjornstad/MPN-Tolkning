from __future__ import annotations

import re
import hashlib
from collections import defaultdict
from collections.abc import Mapping
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Sequence
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.drawing.image import Image as WorksheetImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from mpn_tolkning.core.context import classify_driver_context
from mpn_tolkning.core.identity import stable_variant_key
from mpn_tolkning.core.models import EvidenceResult, VariantRecord
from mpn_tolkning.reports.safety import safe_excel_value


PROVIDER_ORDER = ("CALR-ETdb", "MTBP", "Franklin", "ClinVar", "OncoKB", "COSMIC")
DISCLAIMER = (
    "Research and decision-support output only. Every finding must be reviewed and "
    "approved by qualified laboratory personnel before clinical use."
)
_TEAL = "123C3C"
_ACCENT = "0F766E"
_AMBER = "D97706"
_PALE = "F4F8F7"
_PROVIDER_HOSTS = {
    "CALR-ETdb": {"www.dsimb.inserm.fr"},
    "Franklin": {"franklin.genoox.com"},
    "ClinVar": {"www.ncbi.nlm.nih.gov", "ncbi.nlm.nih.gov"},
    "OncoKB": {"www.oncokb.org", "oncokb.org"},
    "COSMIC": {"cancer.sanger.ac.uk"},
    "ClinGen Allele Registry": {"reg.clinicalgenome.org"},
    "CIViC": {"civicdb.org", "www.civicdb.org"},
    "MetaKB": {"search.cancervariants.org"},
}


def _safe_filename(value: str) -> str:
    readable = re.sub(r"[^A-Za-z0-9._-]+", "_", value).strip("._") or "sample"
    digest = hashlib.sha256(value.encode("utf-8")).hexdigest()[:8]
    return f"{readable}-{digest}"


def _safe_sheet_name(value: str, existing: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "_", value).strip()[:31] or "Variant"
    candidate = base
    counter = 2
    while candidate in existing:
        suffix = f"_{counter}"
        candidate = base[: 31 - len(suffix)] + suffix
        counter += 1
    existing.add(candidate)
    return candidate


def _sorted_evidence(evidence: Sequence[EvidenceResult]) -> list[EvidenceResult]:
    position = {name: index for index, name in enumerate(PROVIDER_ORDER)}
    return sorted(evidence, key=lambda item: (position.get(item.provider, 999), item.provider.lower()))


def _safe_evidence_url(provider: str, value: str) -> bool:
    parsed = urlparse(value)
    return (
        parsed.scheme == "https"
        and parsed.hostname in _PROVIDER_HOSTS.get(provider, set())
    )


def _title(sheet, title: str, subtitle: str = "") -> None:
    sheet.merge_cells("A1:F1")
    sheet["A1"] = title
    sheet["A1"].font = Font(name="Segoe UI", size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=_TEAL)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 32
    if subtitle:
        sheet.merge_cells("A2:F2")
        sheet["A2"] = safe_excel_value(subtitle)
        sheet["A2"].font = Font(name="Segoe UI", size=10, color="3F5656")
        sheet["A2"].fill = PatternFill("solid", fgColor=_PALE)
        sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[2].height = 36


def _style_header(row) -> None:
    for cell in row:
        cell.font = Font(name="Segoe UI", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=_ACCENT)
        cell.alignment = Alignment(vertical="center")


def _configure_sheet(sheet, widths=(18, 24, 28, 18, 18, 42)) -> None:
    sheet.freeze_panes = "A4"
    sheet.sheet_view.showGridLines = False
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.outlinePr.summaryBelow = True


def write_patient_report(
    path: Path,
    sample: str,
    variants: Sequence[VariantRecord],
    evidence: Sequence[EvidenceResult] | Mapping[str, EvidenceResult],
) -> Path:
    if not isinstance(evidence, Mapping) and evidence:
        raise ValueError("Evidence must be keyed by patient, variant, and provider identity.")
    output = Path(path)
    if output.suffix.lower() != ".xlsx":
        output = output / f"{_safe_filename(sample)}_MPN_Tolkning.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    overview = workbook.active
    overview.title = "Oversikt"
    attachments = workbook.create_sheet("Vedlegg")
    _title(overview, "MPN Tolkning", f"Prøve: {sample}")
    _configure_sheet(overview)
    headers = ("Gen", "HGVSc", "HGVSp", "VAF", "Driverkontekst", "Merknad")
    overview.append([])
    overview.append(headers)
    _style_header(overview[4])
    for variant in variants:
        context = classify_driver_context(variant)
        overview.append(
            (
                safe_excel_value(variant.gene),
                safe_excel_value(variant.hgvsc),
                safe_excel_value(variant.hgvsp),
                variant.vaf,
                safe_excel_value(context.label),
                safe_excel_value(context.detail),
            )
        )
        overview.cell(overview.max_row, 2).font = Font(name="Consolas", size=10)
        overview.cell(overview.max_row, 3).font = Font(name="Consolas", size=10)
    disclaimer_row = overview.max_row + 3
    overview.merge_cells(start_row=disclaimer_row, start_column=1, end_row=disclaimer_row, end_column=6)
    overview.cell(disclaimer_row, 1, DISCLAIMER)
    overview.cell(disclaimer_row, 1).font = Font(name="Segoe UI", italic=True, color="7A3E00")
    overview.cell(disclaimer_row, 1).fill = PatternFill("solid", fgColor="FFF4E5")
    overview.cell(disclaimer_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
    overview.row_dimensions[disclaimer_row].height = 42

    _title(attachments, "Vedlegg", "Verifiserte skjermbilder brukt i evidensgjennomgangen")
    _configure_sheet(attachments, (20, 34, 20, 20, 20, 20))
    attachments["A4"] = "Skjermbilder vises også under den relevante varianten."
    attachments["A4"].font = Font(name="Segoe UI", italic=True, color="3F5656")

    existing = set(workbook.sheetnames)
    for index, variant in enumerate(variants, 1):
        label = f"{variant.gene}_{variant.hgvsc.split(':')[-1]}"
        sheet = workbook.create_sheet(_safe_sheet_name(label, existing))
        _title(sheet, f"{variant.gene} variant", f"{variant.hgvsc} · {variant.hgvsp or 'protein consequence not provided'}")
        _configure_sheet(sheet, (20, 22, 52, 48, 20, 20))
        sheet.append([])
        sheet.append(("Database", "Status", "Summary", "Link", "Expected identity", "Observed identity"))
        _style_header(sheet[4])
        row = 5
        if isinstance(evidence, Mapping):
            result_prefix = f"{sample}|{stable_variant_key(variant)}|"
            variant_evidence = [
                result for key, result in evidence.items() if key.startswith(result_prefix)
            ]
        else:
            variant_evidence = []
        for item in _sorted_evidence(variant_evidence):
            sheet.cell(row, 1, safe_excel_value(item.provider))
            sheet.cell(row, 2, item.status.value)
            sheet.cell(row, 3, safe_excel_value(item.summary))
            retained_url = item.url if _safe_evidence_url(item.provider, item.url) else ""
            sheet.cell(row, 4, safe_excel_value(retained_url))
            if retained_url:
                sheet.cell(row, 4).hyperlink = item.url
                sheet.cell(row, 4).style = "Hyperlink"
            sheet.cell(row, 5, safe_excel_value(item.expected_identity))
            sheet.cell(row, 6, safe_excel_value(item.observed_identity))
            for cell in sheet[row]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            row += 1
            for screenshot in item.screenshots:
                if not Path(screenshot).is_file():
                    continue
                sheet.cell(row, 1, f"{item.provider} capture")
                image = WorksheetImage(str(screenshot))
                scale = min(1.0, 900 / max(image.width, 1), 420 / max(image.height, 1))
                image.width = int(image.width * scale)
                image.height = int(image.height * scale)
                image.anchor = f"B{row}"
                sheet.add_image(image)
                sheet.row_dimensions[row].height = max(90, image.height * 0.75)
                row += 1
        if row == 5:
            sheet.cell(row, 1, "No evidence results have been collected yet.")
        sheet.auto_filter.ref = f"A4:F{max(4, row - 1)}"

    thin = Side(style="thin", color="D7E3E1")
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    font = copy(cell.font)
                    font.name = font.name or "Segoe UI"
                    cell.font = font
                    if cell.row >= 4:
                        cell.border = Border(bottom=thin)
    workbook.save(output)
    return output


@dataclass(frozen=True)
class _PendingReport:
    sample: str
    variants: tuple[VariantRecord, ...]
    evidence: Sequence[EvidenceResult] | Mapping[str, EvidenceResult]


class PatientReportCoordinator:
    def __init__(self, output_directory: Path):
        self.output_directory = Path(output_directory)
        self._pending: dict[str, _PendingReport] = {}

    @property
    def pending_samples(self) -> set[str]:
        return set(self._pending)

    def _save(
        self,
        sample: str,
        variants: Sequence[VariantRecord],
        evidence: Sequence[EvidenceResult] | Mapping[str, EvidenceResult],
    ) -> Path:
        return write_patient_report(self.output_directory, sample, variants, evidence)

    def update_patient(
        self,
        sample: str,
        variants: Sequence[VariantRecord],
        evidence: Sequence[EvidenceResult] | Mapping[str, EvidenceResult],
    ) -> Path | None:
        stored_evidence = dict(evidence) if isinstance(evidence, Mapping) else tuple(evidence)
        payload = _PendingReport(sample, tuple(variants), stored_evidence)
        try:
            output = self._save(payload.sample, payload.variants, payload.evidence)
        except PermissionError:
            self._pending[sample] = payload
            return None
        self._pending.pop(sample, None)
        return output

    def reconcile(self) -> dict[str, Path]:
        written: dict[str, Path] = {}
        for sample, payload in list(self._pending.items()):
            try:
                written[sample] = self._save(payload.sample, payload.variants, payload.evidence)
            except PermissionError:
                continue
            self._pending.pop(sample, None)
        return written

    def write_all(
        self,
        variants: Iterable[VariantRecord],
        evidence: Sequence[EvidenceResult] | Mapping[str, EvidenceResult],
    ) -> dict[str, Path]:
        grouped: dict[str, list[VariantRecord]] = defaultdict(list)
        for variant in variants:
            grouped[variant.sample].append(variant)
        written: dict[str, Path] = {}
        for sample, sample_variants in grouped.items():
            if isinstance(evidence, Mapping):
                sample_prefix = f"{sample}|"
                sample_evidence = {
                    key: result for key, result in evidence.items() if key.startswith(sample_prefix)
                }
            else:
                sample_evidence = evidence
            path = self.update_patient(sample, sample_variants, sample_evidence)
            if path is not None:
                written[sample] = path
        return written
